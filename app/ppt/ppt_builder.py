from datetime import datetime
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


class PresentationTheme:
    def __init__(self) -> None:
        self.title_font_name = "Calibri"
        self.body_font_name = "Calibri"
        self.title_font_size = Pt(40)
        self.body_font_size = Pt(24)
        self.small_font_size = Pt(16)


def _add_title_slide(prs: Presentation, session_title: str, presenter_name: Optional[str], affiliation: Optional[str]) -> None:
    slide_layout = prs.slide_layouts[0]  # Title
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = session_title
    subtitle = slide.placeholders[1]

    subtitle_lines: List[str] = []
    if presenter_name:
        subtitle_lines.append(presenter_name)
    if affiliation:
        subtitle_lines.append(affiliation)
    subtitle_lines.append(datetime.now().strftime("%b %d, %Y"))

    subtitle.text = "\n".join(subtitle_lines)


def _add_content_slide(prs: Presentation, theme: PresentationTheme, title: str, bullets: List[str], speaker_notes: str = "", figure_suggestion: str = "") -> None:
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)

    # Title
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = theme.title_font_size

    # Body
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()

    if bullets:
        for idx, bullet in enumerate(bullets):
            if idx == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = theme.body_font_size
    else:
        p = tf.paragraphs[0]
        p.text = "(content to be added)"
        p.font.size = theme.body_font_size

    # Speaker notes
    if speaker_notes:
        notes_slide = slide.notes_slide
        notes_tf = notes_slide.notes_text_frame
        notes_tf.clear()
        notes_p = notes_tf.paragraphs[0]
        notes_p.text = speaker_notes
        notes_p.font.size = theme.small_font_size

    # Figure suggestion as a small textbox in the bottom-right
    if figure_suggestion:
        left = Inches(8.0)
        top = Inches(6.5)
        width = Inches(2.5)
        height = Inches(1.0)
        textbox = slide.shapes.add_textbox(left, top, width, height)
        tb_tf = textbox.text_frame
        tb_tf.text = f"Figure: {figure_suggestion}"
        tb_tf.paragraphs[0].font.size = Pt(12)
        tb_tf.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _add_references_slide(prs: Presentation, theme: PresentationTheme, references: List[str]) -> None:
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "References"

    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()

    for idx, ref in enumerate(references[:12]):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = ref
        p.level = 0
        p.font.size = theme.small_font_size


def _add_disclaimer_slide(prs: Presentation, theme: PresentationTheme) -> None:
    slide_layout = prs.slide_layouts[5]  # Title Only
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Disclaimer"

    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(8.8)
    height = Inches(4.5)

    textbox = slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = (
        "This teaching material is for educational purposes only and does not constitute medical advice. "
        "Verify recommendations against current institutional policies and published guidelines. "
        "Review and update local protocols, dosages, and device availability."
    )
    p.font.size = theme.body_font_size


def build_powerpoint(content: Dict[str, Any], presenter_name: Optional[str] = None, affiliation: Optional[str] = None, output_path: Optional[str] = None) -> str:
    """
    Build a PPTX file from structured content and return the file path.
    content schema:
      {
        "title": str,
        "slides": [
          {"title": str, "bullets": [str], "speakerNotes": str, "figureSuggestion": str}, ...
        ],
        "references": [str]
      }
    """
    prs = Presentation()
    theme = PresentationTheme()

    session_title = content.get("title", "General Surgery Teaching Session")
    slides: List[Dict[str, Any]] = content.get("slides", [])
    references: List[str] = content.get("references", [])

    _add_title_slide(prs, session_title, presenter_name, affiliation)

    for slide in slides:
        _add_content_slide(
            prs,
            theme,
            title=str(slide.get("title", "Untitled")),
            bullets=[str(b) for b in slide.get("bullets", [])],
            speaker_notes=str(slide.get("speakerNotes", "")),
            figure_suggestion=str(slide.get("figureSuggestion", "")),
        )

    if references:
        _add_references_slide(prs, theme, references)

    _add_disclaimer_slide(prs, theme)

    file_path = output_path or f"/workspace/teaching_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
    prs.save(file_path)
    return file_path